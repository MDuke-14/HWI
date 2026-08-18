"""
Rotas do novo sistema de Férias (Código do Trabalho arts. 237.º–246.º, 264.º)
=============================================================================

Complementa `routes/vacations.py`. Não substitui — mantém retro-compat com
os pedidos existentes (`vacation_requests`) e apenas adiciona:

- `vacation_configs` — config permanente do colaborador (admissão, gozados anteriores)
- `vacation_audit`   — histórico de todas as alterações
- `vacation_mapa`    — snapshots do mapa de férias publicado

Endpoints:
- GET /admin/vacations/config/{user_id}
- PUT /admin/vacations/config/{user_id}
- GET /admin/vacations/breakdown/{user_id}
- GET /admin/vacations/audit/{user_id}
- GET /vacations/breakdown            (para o próprio)
- GET /admin/vacations/mapa?year=
- GET /admin/vacations/mapa/excel?year=
- POST /admin/vacations/mapa/publicar?year=
"""
from __future__ import annotations
import io
import logging
import uuid
from datetime import date, datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse

from database import db
from server import get_current_user, get_current_admin, create_notification
from vacation_engine import (
    calcular_saldo_completo, totais_do_saldo, dias_uteis_no_periodo,
)

router = APIRouter()


# =============================================================================
# Helpers internos
# =============================================================================

async def _get_config(user_id: str) -> dict:
    """Devolve config do user (novo modelo). Se não existir mas houver
    vacation_balances com company_start_date, faz upsert-lazy migrando esse
    campo para preservar histórico."""
    cfg = await db.vacation_configs.find_one({"user_id": user_id}, {"_id": 0})
    if cfg:
        return cfg

    # Migração lazy: se existir company_start_date em vacation_balances, cria config
    legacy = await db.vacation_balances.find_one({"user_id": user_id}, {"_id": 0})
    if legacy and legacy.get("company_start_date"):
        cfg = {
            "user_id": user_id,
            "admissao_date": legacy["company_start_date"],
            "admissao_date_set_at": datetime.now(timezone.utc).isoformat(),
            "admissao_date_set_by": "migration_v1",
            "dias_gozados_anteriores": {},  # {year: int}
            "subsidio_ferias_valor": None,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "updated_at": datetime.now(timezone.utc).isoformat(),
        }
        # Importar do vacation_taken_by_year legacy (opcional, se existir)
        taken_docs = await db.vacation_taken_by_year.find(
            {"user_id": user_id}, {"_id": 0}
        ).to_list(None)
        for t in taken_docs:
            try:
                y = int(t.get("year"))
                d = int(t.get("days_taken") or 0)
                if d > 0:
                    cfg["dias_gozados_anteriores"][str(y)] = d
            except (TypeError, ValueError):
                continue
        await db.vacation_configs.insert_one({**cfg})  # insert fresh dict
        logging.info(f"vacation_configs criada (lazy migration) para user={user_id}")
        return cfg

    # Sem qualquer histórico — devolve config vazia (admissão em falta)
    return {
        "user_id": user_id,
        "admissao_date": None,
        "admissao_date_set_at": None,
        "admissao_date_set_by": None,
        "dias_gozados_anteriores": {},
        "subsidio_ferias_valor": None,
    }


async def _audit(
    user_id: str, action: str, before: dict, after: dict, motivo: str, admin: dict
):
    """Persiste 1 entrada no audit trail (sem apagar dados anteriores)."""
    entry = {
        "id": str(uuid.uuid4()),
        "user_id": user_id,
        "action": action,
        "before": before,
        "after": after,
        "motivo": motivo or "",
        "admin_id": admin.get("sub"),
        "admin_username": admin.get("username"),
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.vacation_audit.insert_one(entry)


async def _fetch_saldo_ctx(user_id: str):
    """Devolve tudo o que o engine precisa para calcular."""
    cfg = await _get_config(user_id)
    approved = await db.vacation_requests.find(
        {"user_id": user_id, "status": "approved"}, {"_id": 0},
    ).to_list(None)
    cancelled = await db.cancelled_vacation_days.find(
        {"user_id": user_id}, {"_id": 0, "date": 1},
    ).to_list(None)
    cancelled_set = {c["date"] for c in cancelled}
    # dias_gozados_anteriores está em string {"year": int} — normalizar
    dga = {}
    for k, v in (cfg.get("dias_gozados_anteriores") or {}).items():
        try:
            dga[int(k)] = int(v or 0)
        except (TypeError, ValueError):
            continue
    return cfg, approved, cancelled_set, dga


def _breakdown_from_ctx(cfg, approved, cancelled_set, dga, include_after=1):
    """Constrói o breakdown legal. Falha se admissao_date não estiver definida."""
    if not cfg.get("admissao_date"):
        return {
            "error": "admissao_date_missing",
            "message": "Data de admissão em falta — configure em /admin › Férias › Gerir.",
        }
    try:
        adm = datetime.strptime(cfg["admissao_date"], "%Y-%m-%d").date()
    except ValueError:
        return {"error": "admissao_date_invalid", "message": "Data de admissão inválida."}
    saldos = calcular_saldo_completo(
        admissao_date=adm,
        approved_requests=approved,
        cancelled_dates=cancelled_set,
        dias_gozados_anteriores=dga,
        include_years_after=include_after,
    )
    totais = totais_do_saldo(saldos)
    return {
        "admissao_date": cfg["admissao_date"],
        "subsidio_ferias_valor": cfg.get("subsidio_ferias_valor"),
        "year_breakdown": saldos,
        "totais": totais,
    }


# =============================================================================
# Admin: GET / PUT config
# =============================================================================

@router.get("/admin/vacations/config/{user_id}")
async def admin_get_vacation_config(user_id: str, current_user: dict = Depends(get_current_admin)):
    user = await db.users.find_one({"id": user_id}, {"_id": 0, "hashed_password": 0})
    if not user:
        raise HTTPException(404, "Utilizador não encontrado")
    cfg = await _get_config(user_id)
    return {"user": user, "config": cfg}


@router.put("/admin/vacations/config/{user_id}")
async def admin_put_vacation_config(
    user_id: str, data: dict, current_user: dict = Depends(get_current_admin),
):
    """Actualiza config permanente do colaborador.

    Body: {
      admissao_date: "YYYY-MM-DD"  (obrigatório, gravado permanentemente)
      dias_gozados_anteriores: { "2023": 15, "2024": 10 }  (opcional)
      subsidio_ferias_valor: number|null  (opcional, meta apenas)
      motivo: str  (obrigatório para alterações)
    }
    """
    if not await db.users.find_one({"id": user_id}, {"_id": 1}):
        raise HTTPException(404, "Utilizador não encontrado")

    motivo = (data.get("motivo") or "").strip()
    adm_new = (data.get("admissao_date") or "").strip()
    if not adm_new:
        raise HTTPException(400, "admissao_date é obrigatório")
    try:
        datetime.strptime(adm_new, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(400, "admissao_date com formato inválido (YYYY-MM-DD)")

    existing = await _get_config(user_id)
    is_change = existing.get("admissao_date") is not None
    if is_change and not motivo:
        raise HTTPException(400, "motivo é obrigatório em alterações")

    # dias_gozados_anteriores — validar
    dga_raw = data.get("dias_gozados_anteriores") or {}
    dga_norm = {}
    for k, v in dga_raw.items():
        try:
            y = int(k)
            n = int(v or 0)
            if n < 0:
                raise ValueError
            dga_norm[str(y)] = n
        except (TypeError, ValueError):
            raise HTTPException(400, f"dias_gozados_anteriores[{k}] inválido")

    now = datetime.now(timezone.utc).isoformat()
    new_doc = {
        "user_id": user_id,
        "admissao_date": adm_new,
        "admissao_date_set_at": existing.get("admissao_date_set_at") or now,
        "admissao_date_set_by": existing.get("admissao_date_set_by") or current_user.get("username"),
        "dias_gozados_anteriores": dga_norm,
        "subsidio_ferias_valor": data.get("subsidio_ferias_valor"),
        "created_at": existing.get("created_at") or now,
        "updated_at": now,
    }
    await db.vacation_configs.update_one(
        {"user_id": user_id}, {"$set": new_doc}, upsert=True,
    )

    # Retro-compat: também guarda em vacation_balances.company_start_date
    # (para não partir código legado enquanto migração completa não acaba)
    await db.vacation_balances.update_one(
        {"user_id": user_id},
        {"$set": {"user_id": user_id, "company_start_date": adm_new}},
        upsert=True,
    )

    # Audit
    await _audit(
        user_id,
        action="update_config",
        before={k: existing.get(k) for k in ("admissao_date", "dias_gozados_anteriores", "subsidio_ferias_valor")},
        after={k: new_doc.get(k) for k in ("admissao_date", "dias_gozados_anteriores", "subsidio_ferias_valor")},
        motivo=motivo,
        admin=current_user,
    )
    return {"success": True, "config": new_doc}


# =============================================================================
# Breakdown (Admin + próprio)
# =============================================================================

@router.get("/admin/vacations/breakdown/{user_id}")
async def admin_get_breakdown(
    user_id: str,
    include_years_after: int = Query(1, ge=0, le=3),
    current_user: dict = Depends(get_current_admin),
):
    cfg, approved, cancelled_set, dga = await _fetch_saldo_ctx(user_id)
    result = _breakdown_from_ctx(cfg, approved, cancelled_set, dga, include_years_after)
    result["user_id"] = user_id
    return result


@router.get("/vacations/breakdown")
async def user_get_breakdown(current_user: dict = Depends(get_current_user)):
    cfg, approved, cancelled_set, dga = await _fetch_saldo_ctx(current_user["sub"])
    result = _breakdown_from_ctx(cfg, approved, cancelled_set, dga, include_after=1)
    result["user_id"] = current_user["sub"]
    return result


@router.get("/admin/vacations/audit/{user_id}")
async def admin_get_audit(
    user_id: str,
    limit: int = Query(100, ge=1, le=500),
    current_user: dict = Depends(get_current_admin),
):
    entries = await db.vacation_audit.find(
        {"user_id": user_id}, {"_id": 0},
    ).sort("created_at", -1).limit(limit).to_list(limit)
    return {"entries": entries}


# =============================================================================
# Mapa de Férias
# =============================================================================

@router.get("/admin/vacations/mapa")
async def admin_get_mapa(
    year: int = Query(...),
    current_user: dict = Depends(get_current_admin),
):
    """Devolve o mapa de férias do ano (colaborador × períodos)."""
    users = await db.users.find(
        {"is_active": True}, {"_id": 0, "hashed_password": 0},
    ).to_list(1000)
    rows = []
    for u in users:
        approved = await db.vacation_requests.find(
            {"user_id": u["id"], "status": "approved"}, {"_id": 0},
        ).to_list(None)
        # filtra períodos que se intersectam com o ano
        periodos = []
        for req in approved:
            try:
                s = datetime.strptime(req["start_date"], "%Y-%m-%d").date()
                e = datetime.strptime(req["end_date"], "%Y-%m-%d").date()
            except (KeyError, ValueError):
                continue
            if e.year < year or s.year > year:
                continue
            periodos.append({
                "start": max(s, date(year, 1, 1)).isoformat(),
                "end": min(e, date(year, 12, 31)).isoformat(),
                "days_uteis": dias_uteis_no_periodo(
                    max(s, date(year, 1, 1)),
                    min(e, date(year, 12, 31)),
                ),
                "status": req.get("status"),
            })
        # Ordenar por start
        periodos.sort(key=lambda p: p["start"])
        rows.append({
            "user_id": u["id"],
            "username": u.get("username"),
            "full_name": u.get("full_name") or u.get("username"),
            "periodos": periodos,
            "total_dias": sum(p["days_uteis"] for p in periodos),
        })
    # Snapshot publicado (se existir)
    publicado = await db.vacation_mapa.find_one({"year": year}, {"_id": 0})
    return {
        "year": year,
        "rows": rows,
        "publicado": publicado,
    }


@router.get("/admin/vacations/mapa/excel")
async def admin_get_mapa_excel(
    year: int = Query(...),
    current_user: dict = Depends(get_current_admin),
):
    """Exporta o mapa em Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    mapa = await admin_get_mapa(year=year, current_user=current_user)  # reuse

    wb = Workbook()
    ws = wb.active
    ws.title = f"Mapa Ferias {year}"

    ws["A1"] = f"Mapa de Férias — {year}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:E1")

    headers = ["Colaborador", "Início", "Fim", "Dias Úteis", "Estado"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="333333")
        c.alignment = Alignment(horizontal="center")

    row = 4
    for r in mapa["rows"]:
        if not r["periodos"]:
            ws.cell(row=row, column=1, value=r["full_name"])
            ws.cell(row=row, column=5, value="Sem férias marcadas")
            row += 1
            continue
        for p in r["periodos"]:
            ws.cell(row=row, column=1, value=r["full_name"])
            ws.cell(row=row, column=2, value=p["start"])
            ws.cell(row=row, column=3, value=p["end"])
            ws.cell(row=row, column=4, value=p["days_uteis"])
            ws.cell(row=row, column=5, value=p["status"])
            row += 1

    for col_letter, w in zip("ABCDE", (28, 14, 14, 12, 14)):
        ws.column_dimensions[col_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"Mapa_Ferias_{year}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/admin/vacations/mapa/publicar")
async def admin_publicar_mapa(
    year: int = Query(...),
    current_user: dict = Depends(get_current_admin),
):
    """Publica um snapshot do mapa de férias do ano.

    Cria um documento em `vacation_mapa` com o snapshot actual + notifica
    todos os colaboradores activos. Deadline legal: 15 Abril.
    """
    mapa = await admin_get_mapa(year=year, current_user=current_user)
    snapshot = {
        "year": year,
        "published_at": datetime.now(timezone.utc).isoformat(),
        "published_by_id": current_user.get("sub"),
        "published_by_name": current_user.get("username"),
        "rows": mapa["rows"],
    }
    await db.vacation_mapa.update_one(
        {"year": year}, {"$set": snapshot}, upsert=True,
    )
    # Notifica utilizadores
    users = await db.users.find(
        {"is_active": True}, {"_id": 0, "id": 1, "full_name": 1, "username": 1},
    ).to_list(1000)
    for u in users:
        try:
            await create_notification(
                u["id"], "mapa_ferias_publicado",
                f"O mapa de férias de {year} foi publicado. Consulte em /vacations.",
                None,
            )
        except Exception as e:
            logging.warning(f"Falha ao notificar mapa a {u.get('username')}: {e}")
    return {"success": True, "year": year, "notified": len(users)}
